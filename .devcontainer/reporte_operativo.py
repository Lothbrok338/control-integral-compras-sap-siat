"""Reporte operativo compacto para Control Integral de Compras SAP–SIAT.

NO modifica el motor_auditoria.py ni la clasificación de facturas.
Se ejecuta sobre el Excel que produce run_audit() y el Excel SAP ORIGINAL.

Pestañas producidas:
  1. CONTROL GENERAL (con usuario SAP y asiento en C/D)
  2. POR CONCILIAR (solo SAP, solo SIAT, candidatos)
  3. RESUMEN (indicadores operativos)

Por seguridad, nunca se atribuye usuario/asiento usando únicamente similitud:
se cotejan el NIT y la factura, más la autorización cuando está disponible.
Si el origen es ambiguo, se deja en blanco y se notifica la cobertura.
"""
from __future__ import annotations

import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAGENTA = 'C5004E'
GRAFITO = '31343A'
BLANCO = 'FFFFFF'
TEXTO = '343A40'
GRIS = '667085'
BORDE = 'DFE3E8'
FONDO = 'F7F8FA'
MAGENTA_SUAVE = 'FCE8F0'
AMARILLO_SUAVE = 'FFF3CE'
VERDE_SUAVE = 'E8F6EF'
ROJO_SUAVE = 'FDECEB'
AZUL_SUAVE = 'ECF3FB'
NUM_FORMAT = '#,##0.00;[Red](#,##0.00);–'


def _norm(text: Any) -> str:
    if text is None:
        return ''
    t = unicodedata.normalize('NFKD', str(text))
    t = ''.join(ch for ch in t if not unicodedata.combining(ch))
    return re.sub(r'[^A-Z0-9]+', ' ', t.upper()).strip()


def _ident(value: Any) -> str:
    """Normaliza identificadores para cotejo (sin inventar equivalencias)."""
    if value is None or str(value).strip() == '':
        return ''
    v = str(value).strip()
    # Algunos lectores convierten el número de factura 123 en 123.0.
    if re.fullmatch(r'\d+\.0+', v):
        v = v.split('.')[0]
    # 00123 y 123 son equivalentes solo para valores estrictamente numéricos.
    if re.fullmatch(r'\d+', v):
        return v.lstrip('0') or '0'
    return _norm(v).replace(' ', '')


def _auth(value: Any) -> str:
    return _norm(value).replace(' ', '')


def _read_sheet(wb, name: str) -> list[dict]:
    if name not in wb.sheetnames:
        return []
    rows = wb[name].iter_rows(values_only=True)
    try:
        headers = [str(x).strip() if x is not None else '' for x in next(rows)]
    except StopIteration:
        return []
    return [dict(zip(headers, vals)) for vals in rows if any(x is not None for x in vals)]


def _safe(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value  # Evita que datos externos se interpreten como fórmulas.
    return value


def _key(fila: Any) -> int | None:
    try:
        return int(fila)
    except (ValueError, TypeError):
        return None


def _classify_header(txt: Any) -> str | None:
    t = _norm(txt)
    if not t:
