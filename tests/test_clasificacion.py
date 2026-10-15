"""Clasificación crítica / menor de diferencias (datos SINTÉTICOS)."""
from __future__ import annotations

import json

import openpyxl
import pytest

import clasificacion as cl
import reporte_operativo as ro
from motor_auditoria import run_audit
from sinteticos import escenario, escribir_sap, escribir_siat, sap_row, siat_row


# ------------------------------------------------------------- razón social
@pytest.mark.parametrize('sap,siat,motivo', [
    ('INTERCOM LTDA.', 'INTERCOM L.T.D.A.', 'formato'),
    ('CRISIL SRL.', 'CRISIL S.R.L.', 'formato'),
    ('MARCELO PINTO MACEDO', 'PINTO MACEDO MARCELO', 'orden'),
    ('STHEPHANY INES MORRIS FLORES', 'STEPHANY INES MORRIS FLORES', 'ortográfica'),
    ('CERTIFICACIONES DIGITALES DIGICERTS.R.L.', 'CERTIFICACIONES DIGITALES DIGICERT S.R.L.', 'formato'),
    ('CIA ANDINA SA', 'COMPAÑIA ANDINA S.A.', 'Abreviatura'),
    ('JOSÉ  PÉREZ', 'JOSE PEREZ', 'formato'),
    ('SOCIEDAD ANONIMA DEL SUR', 'S.A. SUR', 'formato'),
])
def test_variaciones_menores_con_nit_consistente(sap, siat, motivo):
    nivel, texto = cl.clasificar_razon_social(sap, siat, '123')
    assert nivel == 'MENOR', texto
    assert motivo.lower() in texto.lower()


@pytest.mark.parametrize('sap,siat', [
    ('JUAN PEREZ', 'MARIA LOPEZ'),                                   # nombres distintos
    ('PEDRO RICHARD AQUINO ENCINAS', 'FERRETERIA URKUPINA DE AQUINO ENCINAS PEDRO RICHARD'),  # persona vs comercio
    ('UNIVALLE S.A.', 'UNIVERSIDAD PRIVADA DEL VALLE S.A.'),         # sigla sin equivalencia verificada
    ('PEDRO RUIZ SRL', 'PEDRO RUIZ LTDA'),                           # forma jurídica distinta
    ('PEDRO RUIZ', 'PEDRO RUIZ SRL'),                                # forma jurídica omitida
    ('CAJA PETROLERA DE SALUD DEPTAL STA CRUZ', 'CAJA PETROLERA DE SALUD'),  # un nombre contiene al otro
    ('JUAN PEREZ', 'JUANA PEREZ'),                                   # similitud alta, token corto: no se aprueba
    ('', 'ALGO SA'),
])
def test_casos_que_requieren_revision_no_se_aprueban_por_similitud(sap, siat):
    nivel, _ = cl.clasificar_razon_social(sap, siat, '123')
    assert nivel == 'REVISION'


def test_equivalencia_verificada_solo_aplica_a_su_nit(tmp_path):
    eq = [{'nit': '555', 'nombres': ['SIGLA PRUEBA', 'SOCIEDAD DE PRUEBA INTEGRAL'], 'nota': 'verificada'}]
    assert cl.clasificar_razon_social('SIGLA PRUEBA', 'SOCIEDAD DE PRUEBA INTEGRAL', '555', eq)[0] == 'MENOR'
    assert cl.clasificar_razon_social('SIGLA PRUEBA', 'SOCIEDAD DE PRUEBA INTEGRAL', '999', eq)[0] == 'REVISION'
    assert cl.clasificar_razon_social('SIGLA PRUEBA', 'OTRA SOCIEDAD', '555', eq)[0] == 'REVISION'
    ruta = tmp_path / 'eq.json'
    ruta.write_text(json.dumps({'equivalencias': eq}), encoding='utf-8')
    assert cl.cargar_equivalencias(ruta) == eq
    assert cl.cargar_equivalencias(tmp_path / 'no_existe.json') == []
    ruta.write_text('{no es json', encoding='utf-8')
    assert cl.cargar_equivalencias(ruta) == []


def test_archivo_de_equivalencias_del_repositorio_no_trae_nada_inventado():
    assert cl.cargar_equivalencias() == []


# --------------------------------------------- integración con el motor real
def _corrida(tmp_path, pares, eq=None):
    """pares: lista de (sap_row_kwargs, siat_row_kwargs) con factura/NIT únicos."""
    sap_rows, siat_rows = [], []
    for i, (s, t) in enumerate(pares):
        sap_rows.append(sap_row(**s))
        siat_rows.append(siat_row(i + 1, **t))
    sap = escribir_sap(tmp_path / 's.xlsx', sap_rows)
    siat = escribir_siat(tmp_path / 't.csv', siat_rows)
    motor = tmp_path / 'm.xlsx'
    _, m = run_audit(sap, siat, output_path=motor)
    eq_path = None
    if eq is not None:
        eq_path = tmp_path / 'eq.json'
        eq_path.write_text(json.dumps({'equivalencias': eq}), encoding='utf-8')
    diag = ro.generar_reporte_operativo(motor, sap, tmp_path / 'r.xlsx', tolerancia=0.05,
                                        equivalencias_path=eq_path)
    wb = openpyxl.load_workbook(tmp_path / 'r.xlsx')

    def filas(h):
        ws = wb[h]
        cab = [c.value for c in ws[1]]
        return {r[cab.index('N.º FACTURA')]: dict(zip(cab, r)) for r in ws.iter_rows(min_row=2, values_only=True)}

    return filas('CONTROL GENERAL'), filas('DIFERENCIAS MENORES'), m, diag


def _par(n, razon_sap, razon_siat, nit_sap=None, nit_siat=None, total=100.0, total_siat=None, fecha_siat=None, over=None):
    nit = f'30000000{n:02d}'
    aut = f'AUTORIZ{n:08d}'
    fecha = '2025-02-10'
    return (dict(nit=nit_sap or nit, factura=n, aut=aut, fecha=fecha, total=total, razon=razon_sap),
            dict(nit=nit_siat or nit, factura=n, aut=aut, fecha=fecha_siat or fecha,
                 total=total_siat if total_siat is not None else total, razon=razon_siat, over=over))


def test_clasificacion_integrada_con_el_motor(tmp_path):
    pares = [
        _par(1, 'INTERCOM LTDA.', 'INTERCOM L.T.D.A.'),                                   # menor
        _par(2, 'MARCELO PINTO MACEDO', 'PINTO MACEDO MARCELO'),                          # menor
        _par(3, 'EMPRESA UNO SRL', 'COMERCIAL OTRA SRL'),                                 # revisión de identidad
        _par(4, 'CRISIL SRL.', 'CRISIL S.R.L.', nit_siat='3999999999'),                  # NIT distinto + nombre formal -> crítica
        _par(5, 'ACME LTDA.', 'ACME L.T.D.A.', total=100.0, total_siat=130.0),            # menor + importe -> crítica
        _par(6, 'BETA SA', 'BETA S.A.', total=100.0, total_siat=100.03),                  # menor + redondeo
        _par(7, 'GAMA SA', 'GAMA SA', total=100.0, total_siat=100.03),                    # solo redondeo
        _par(8, 'DELTA SA', 'DELTA SA', fecha_siat='2025-02-11', over={18: '87.03'}),     # fecha crítica + redondeo
        _par(9, 'EPSILON SA', 'EPSILON SA'),                                              # OK
    ]
    control, menores, m, diag = _corrida(tmp_path, pares)

    # --- destino de cada factura
    assert set(control) == {'3', '4', '5', '8'}
    assert set(menores) == {'1', '2', '6', '7'}
    assert '9' not in control and '9' not in menores                     # OK no se exporta
    assert not set(control) & set(menores)                                # nunca en ambas pestañas

    # --- NIT diferente: siempre crítica prioritaria, aunque el nombre solo difiera en puntuación
    assert control['4']['ESTADO'] == 'CRÍTICA · IDENTIFICADOR'
    assert 'NIT' in control['4']['CAMPOS CON DIFERENCIA']
    assert 'Razón social: REVISIÓN DE IDENTIDAD' in control['4']['CLASIFICACIÓN POR CAMPO']  # no se aprueba por puntuación

    # --- una menor no oculta una crítica
    assert control['5']['ESTADO'] == 'CRÍTICA · IMPORTE O TRIBUTARIO'
    assert 'Razón social: MENOR' in control['5']['CLASIFICACIÓN POR CAMPO']
    assert control['8']['ESTADO'] == 'CRÍTICA · FECHA'
    assert 'redondeo' in control['8']['VALOR EN SAP']                    # el redondeo acompaña a la crítica

    # --- identidad dudosa permanece en CONTROL GENERAL
    assert control['3']['ESTADO'] == 'REVISIÓN DE IDENTIDAD'
    assert control['3']['RAZÓN SOCIAL SIAT'] == 'COMERCIAL OTRA SRL'

    # --- menores con su motivo
    assert menores['1']['TIPO DE DIFERENCIA'] == 'RAZÓN SOCIAL'
    assert menores['2']['MOTIVO DE CLASIFICACIÓN MENOR'] == 'Cambio de orden de nombres y apellidos'
    assert menores['6']['TIPO DE DIFERENCIA'] == 'RAZÓN SOCIAL + REDONDEO'
    assert menores['7']['TIPO DE DIFERENCIA'] == 'REDONDEO'
    assert all(r['MOTIVO DE CLASIFICACIÓN MENOR'] for r in menores.values())

    # --- los valores monetarios del redondeo se conservan exactos
    assert (menores['7']['IMPORTE SAP (Bs)'], menores['7']['IMPORTE SIAT (Bs)']) == (100, 100.03)

    # --- indicadores: agrupación sin perder facturas; cifras del motor intactas
    ind = diag['indicadores']
    assert (ind['criticas'], ind['revision_identidad'], ind['menores_nominales'],
            ind['solo_redondeo'], ind['sin_diferencias']) == (3, 1, 3, 1, 1)
    assert m['matched'] == 9 and m['ok'] == 1


def test_todo_diferente_de_razon_social_con_nit_consistente_usa_equivalencia_verificada(tmp_path):
    pares = [_par(1, 'SIGLA PRUEBA', 'SOCIEDAD DE PRUEBA INTEGRAL')]
    nit = pares[0][0]['nit']
    control, menores, *_ = _corrida(tmp_path, pares)
    assert set(control) == {'1'} and not menores                         # sin verificación: revisión
    control, menores, *_ = _corrida(tmp_path, pares, eq=[{'nit': nit, 'nombres': ['SIGLA PRUEBA', 'SOCIEDAD DE PRUEBA INTEGRAL']}])
    assert set(menores) == {'1'} and not control                         # verificada para ese NIT


def test_invariante_ninguna_critica_en_menores(tmp_path):
    """Toda diferencia de un campo distinto de razón social/redondeo deja la factura en CONTROL GENERAL."""
    sap, siat = escenario(tmp_path)
    motor = tmp_path / 'm.xlsx'
    run_audit(sap, siat, output_path=motor)
    ro.generar_reporte_operativo(motor, sap, tmp_path / 'r.xlsx')
    mw = openpyxl.load_workbook(motor)
    cab = [c.value for c in mw['DETALLE DIFERENCIAS'][1]]
    no_razon = {r[cab.index('FILA SAP')] for r in mw['DETALLE DIFERENCIAS'].iter_rows(min_row=2, values_only=True)
                if r[cab.index('CAMPO OBSERVADO')] != 'Razón social'}
    wb = openpyxl.load_workbook(tmp_path / 'r.xlsx')
    ws = wb['DIFERENCIAS MENORES']
    c2 = [c.value for c in ws[1]]
    en_menores = {r[c2.index('FILA SAP ORIGINAL')] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert not no_razon & en_menores


def test_autorizacion_o_factura_distinta_con_mismo_nit_no_deja_aprobar_el_nombre(tmp_path):
    """Mismo NIT y nombre con solo puntuación, pero la autorización/CUF difiere: sigue siendo crítica."""
    pares = [_par(1, 'CRISIL SRL.', 'CRISIL S.R.L.', over={3: 'OTRA-AUTORIZACION-0001'})]
    control, menores, *_ = _corrida(tmp_path, pares)
    assert set(control) == {'1'} and not menores
    assert control['1']['ESTADO'] == 'CRÍTICA · IDENTIFICADOR'
    assert 'Código de autorización' in control['1']['CAMPOS CON DIFERENCIA']
    assert 'Razón social: REVISIÓN DE IDENTIDAD' in control['1']['CLASIFICACIÓN POR CAMPO']
