"""Reporte operativo compacto para Control Integral de Compras SAP–SIAT.

Capa POSTERIOR a ``run_audit()``: no modifica ``motor_auditoria.py`` ni el
emparejamiento. Lee el Excel que genera el motor y el Excel SAP ORIGINAL cargado
por el usuario, clasifica las diferencias por gravedad (``clasificacion.py``) y
produce un Excel de solo tres pestañas:

  1. CONTROL GENERAL       solo diferencias críticas o de identidad (A=ESTADO,
                           B=CAMPOS CON DIFERENCIA, C=USUARIO SAP, D=N.º ASIENTO SAP)
  2. POR CONCILIAR         sin pareja (solo SAP / solo SIAT) y posibles emparejamientos
  3. DIFERENCIAS MENORES   informativas: exclusivamente razón social o exclusivamente redondeo

Usuario y asiento SAP solo se asignan con evidencia verificable: la fila de
origen del motor (``_FILA_ORIGEN``) debe coincidir en NIT + factura (+ código de
autorización si ambos lo traen). Si no hay evidencia inequívoca la celda queda
vacía y se informa una advertencia. Nunca se infiere por similitud.
"""
from __future__ import annotations

import math
import numbers
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from clasificacion import (
    CRIT_IDENT, RANGO_ESTADO, REV_IDENTIDAD, TAG_CRITICA, TAG_IDENTIDAD, TAG_MENOR,
    cargar_equivalencias, clasificar_facturas, tipo_menor,
)
from motor_auditoria import (
    SAP_TO_CANON,
    __author__,
    __project__,
    __version__,
    norm_code,
    norm_numeric_id,
)

# ---------------------------------------------------------------------------
# Identidad visual UNIVALLE: magenta + grafito + fondo blanco
# ---------------------------------------------------------------------------
MAGENTA = 'C5004E'
GRAFITO = '31343A'
BLANCO = 'FFFFFF'
TEXTO = '343A40'
GRIS = '667085'
BORDE = 'DFE3E8'
FONDO = 'F7F8FA'
AMARILLO_SUAVE = 'FFF3CE'
VERDE_SUAVE = 'E8F6EF'
ROJO_SUAVE = 'FDECEB'
AZUL_SUAVE = 'ECF3FB'
MAGENTA_SUAVE = 'FCE8F0'
NUM_FORMAT = '#,##0.00;[Red]-#,##0.00'
FUENTE = 'Calibri'

HOJA_CONTROL = 'CONTROL GENERAL'
HOJA_PENDIENTES = 'POR CONCILIAR'
HOJA_MENORES = 'DIFERENCIAS MENORES'
HOJAS_MOTOR_REQUERIDAS = (
    'CONTROL GENERAL', 'DETALLE DIFERENCIAS', 'REDONDEOS',
    'SOLO SAP', 'SOLO SIAT', 'BASE SAP NORMALIZADA',
)
AMARILLO_PESTANA = 'FFE699'

# Encabezados SAP que el motor ya exige (SAP_TO_CANON): se usan tal cual para
# NIT, factura y autorización. Solo usuario y asiento se detectan por alias.
_COL_NIT = next(k for k, v in SAP_TO_CANON.items() if v == 'NIT')
_COL_FACTURA = next(k for k, v in SAP_TO_CANON.items() if v == 'FACTURA')
_COL_AUTORIZACION = next(k for k, v in SAP_TO_CANON.items() if v == 'AUTORIZACION')


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def _norm(text: Any) -> str:
    """Mayúsculas, sin tildes ni signos: 'N.º Asiento' -> 'N O ASIENTO'."""
    if text is None:
        return ''
    t = unicodedata.normalize('NFKD', str(text))
    t = ''.join(ch for ch in t if not unicodedata.combining(ch))
    return re.sub(r'[^A-Z0-9]+', ' ', t.upper()).strip()


def _vacio(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    try:
        if value is pd.NA or value is pd.NaT:
            return True
    except Exception:  # pragma: no cover
        pass
    return str(value).strip() == ''


def _texto_id(value: Any) -> str:
    """Convierte un identificador a texto SIN perder ceros a la izquierda.

    Los textos se conservan tal cual; los enteros/flotantes enteros que el
    lector convirtió a número se devuelven sin '.0'.
    """
    if _vacio(value):
        return ''
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, numbers.Integral):
        return str(int(value))
    if isinstance(value, numbers.Real):
        f = float(value)
        return str(int(f)) if f.is_integer() else ('%.15g' % f)
    s = str(value).strip()
    if re.fullmatch(r'\d+\.0+', s):  # '123.0' producido por algún lector
        s = s.split('.')[0]
    return s


def _key(fila: Any) -> int | None:
    try:
        return int(fila)
    except (ValueError, TypeError):
        return None


def _num(value: Any) -> float:
    try:
        return float(value)
    except (ValueError, TypeError):
        return 0.0


def _read_sheet(wb, name: str) -> list[dict]:
    if name not in wb.sheetnames:
        return []
    rows = wb[name].iter_rows(values_only=True)
    try:
        headers = [str(x).strip() if x is not None else '' for x in next(rows)]
    except StopIteration:
        return []
    return [dict(zip(headers, vals)) for vals in rows if any(x is not None for x in vals)]


# ---------------------------------------------------------------------------
# Detección conservadora de columnas de usuario y asiento en el SAP original
# ---------------------------------------------------------------------------
_USUARIO_TOKENS = {'USUARIO', 'USER', 'USERNAME', 'ERNAM', 'USNAM', 'CREADOR'}
_USUARIO_FRASES = ('CREADO POR', 'REGISTRADO POR', 'INGRESADO POR', 'CAPTURADO POR')
_USUARIO_PREFERIDO = {'CREADO', 'CREACION', 'REGISTRO', 'REGISTRADO', 'INGRESO',
                      'INGRESADO', 'CAPTURA', 'CAPTURADO', 'CREADOR', 'ERNAM'}
# Usuario que modificó/aprobó/autorizó NO es quien registró el documento.
_USUARIO_EXCLUIR_PREFIJOS = ('MODIF', 'CAMBI', 'ACTUALIZ', 'APROB', 'AUTORIZ', 'ANUL')
_ASIENTO_TOKENS = {'ASIENTO', 'BELNR'}
_ASIENTO_FRASES = ('DOCUMENTO CONTABLE', 'DOC CONTABLE', 'DOCUMENTO SAP', 'DOC SAP')
# Encabezado genérico, válido SOLO como coincidencia exacta. Confirmado en la exportación SAP
# real (septiembre 2026): 'Nro. Documento' = documento contable de 10 dígitos, clases KR/RE.
_ASIENTO_EXACTOS = {'NRO DOCUMENTO', 'NRO DE DOCUMENTO', 'NUMERO DOCUMENTO', 'NUMERO DE DOCUMENTO'}
# Un asiento nunca es un dato de la factura.
_ASIENTO_EXCLUIR = {'FACTURA', 'DUI', 'DIM', 'AUTORIZACION', 'CUF', 'NIT', 'CONTROL'}


def _classify_header(txt: Any) -> str | None:
    """Devuelve 'usuario', 'asiento' o None. Coincidencia por palabras completas."""
    t = _norm(txt)
    if not t:
        return None
    tokens = set(t.split())
    marco = f' {t} '

    es_asiento = (bool(tokens & _ASIENTO_TOKENS) or t in _ASIENTO_EXACTOS
                  or any(f' {f} ' in marco for f in _ASIENTO_FRASES))
    if es_asiento and not (tokens & _ASIENTO_EXCLUIR):
        return 'asiento'

    es_usuario = bool(tokens & _USUARIO_TOKENS) or any(f' {f} ' in marco for f in _USUARIO_FRASES)
    if es_usuario and not any(tok.startswith(_USUARIO_EXCLUIR_PREFIJOS) for tok in tokens):
        return 'usuario'
    return None


def detectar_columnas_sap(columnas: list[str]) -> tuple[dict[str, str | None], list[str]]:
    """Elige la columna de usuario y la de asiento; si hay duda, ninguna.

    Devuelve ({'usuario': nombre|None, 'asiento': nombre|None}, advertencias).
    """
    candidatos: dict[str, list[str]] = {'usuario': [], 'asiento': []}
    for col in columnas:
        rol = _classify_header(col)
        if rol:
            candidatos[rol].append(col)

    elegidas: dict[str, str | None] = {'usuario': None, 'asiento': None}
    avisos: list[str] = []
    for rol, lista in candidatos.items():
        if len(lista) == 1:
            elegidas[rol] = lista[0]
        elif len(lista) > 1:
            pref = [c for c in lista if set(_norm(c).split()) & _USUARIO_PREFERIDO] if rol == 'usuario' else []
            if len(pref) == 1:
                elegidas[rol] = pref[0]
            else:
                avisos.append(
                    f"Hay varias columnas SAP que podrían ser el {rol} ({', '.join(map(str, lista))}). "
                    f"No se eligió ninguna para evitar datos incorrectos: la columna de {rol} queda vacía."
                )
    return elegidas, avisos


# ---------------------------------------------------------------------------
# Vinculación verificable con el Excel SAP original
# ---------------------------------------------------------------------------
def _diag_vacio(advertencias: list[str], total: int = 0, **extra) -> dict:
    d = {
        'usuario_detectado': False, 'asiento_detectado': False,
        'columna_usuario': None, 'columna_asiento': None,
        'total_sap': total, 'filas_vinculadas': 0,
        'filas_con_usuario': 0, 'filas_con_asiento': 0,
        'vinculo_fila': 0, 'vinculo_clave': 0, 'vinculo_duplicado_identico': 0,
        'ambiguas': 0, 'sin_coincidencia': 0, 'no_verificables': 0,
        'advertencias': advertencias,
    }
    d.update(extra)
    return d


def _vincular_usuario_asiento(
    sap_original: Path,
    base_sap: list[dict],
    columna_usuario: str | None = None,
    columna_asiento: str | None = None,
) -> tuple[dict[int, tuple[str | None, str | None]], dict]:
    """Devuelve ({fila_origen: (usuario, asiento)}, diagnóstico).

    ``_FILA_ORIGEN`` del motor = posición en ``pd.read_excel(sap)`` + 2, por lo
    que se lee el SAP original con la misma llamada (primera hoja, encabezado en
    la fila 1) y se usa esa posición como pista, SIEMPRE verificada por
    NIT + factura (+ autorización).
    """
    filas_motor = {}
    for rec in base_sap:
        fila = _key(rec.get('_FILA_ORIGEN'))
        if fila is not None:
            filas_motor[fila] = rec
    if not filas_motor:
        return {}, _diag_vacio(['No hay filas SAP normalizadas para enlazar con el archivo original.'])

    total = len(filas_motor)
    try:
        src = pd.read_excel(sap_original, dtype=object)
    except Exception as exc:  # archivo ilegible: el motor ya lo habría rechazado
        return {}, _diag_vacio(
            [f'No se pudo releer el Excel SAP original para recuperar usuario y asiento ({type(exc).__name__}).'],
            total)

    columnas = list(src.columns)
    faltan = [c for c in (_COL_NIT, _COL_FACTURA) if c not in columnas]
    if faltan:
        return {}, _diag_vacio(
            ['El Excel SAP no contiene las columnas ' + ', '.join(faltan) + '; no se puede verificar el enlace.'],
            total)

    avisos: list[str] = []
    detectadas, avisos_det = detectar_columnas_sap([str(c) for c in columnas])
    avisos += avisos_det
    ucol = columna_usuario if columna_usuario in columnas else detectadas['usuario']
    acol = columna_asiento if columna_asiento in columnas else detectadas['asiento']

    # Seguridad: un asiento jamás debe ser el número de factura.
    if acol is not None:
        iguales = sum(
            _texto_id(a) != '' and _texto_id(a) == _texto_id(f)
            for a, f in zip(src[acol], src[_COL_FACTURA])
        )
        if len(src) and iguales / len(src) >= 0.9:
            avisos.append(
                f"La columna «{acol}» coincide con el número de factura en casi todas las filas; "
                'se descartó como asiento para no confundir ambos datos.'
            )
            acol = None

    if ucol is None:
        avisos.append('No se encontró una columna de usuario (registrado por) en la hoja SAP: USUARIO SAP queda vacío.')
    if acol is None:
        avisos.append('No se encontró una columna de asiento/documento contable en la hoja SAP: N.º ASIENTO SAP queda vacío.')
    if ucol is None or acol is None:
        encabezados = ', '.join(str(c) for c in columnas[:60])
        avisos.append(f'Encabezados de la hoja SAP leída: {encabezados}')

    if ucol is None and acol is None:
        return {}, _diag_vacio(avisos, total)

    s_nit = [norm_numeric_id(v) for v in src[_COL_NIT]]
    s_fac = [norm_numeric_id(v) for v in src[_COL_FACTURA]]
    s_aut = [norm_code(v) for v in src[_COL_AUTORIZACION]] if _COL_AUTORIZACION in columnas else [''] * len(src)
    s_usr = [_texto_id(v) for v in src[ucol]] if ucol is not None else [''] * len(src)
    s_asi = [_texto_id(v) for v in src[acol]] if acol is not None else [''] * len(src)

    por_clave: dict[tuple[str, str], list[int]] = defaultdict(list)
    for pos, (n, f) in enumerate(zip(s_nit, s_fac)):
        if n and f:
            por_clave[(n, f)].append(pos)

    def compatible(pos: int, nit: str, fac: str, aut: str) -> bool:
        if s_nit[pos] != nit or s_fac[pos] != fac:
            return False
        return not (aut and s_aut[pos] and aut != s_aut[pos])  # autorización contradictoria

    meta: dict[int, tuple[str | None, str | None]] = {}
    d = _diag_vacio(avisos, total, columna_usuario=ucol, columna_asiento=acol,
                    usuario_detectado=ucol is not None, asiento_detectado=acol is not None)
    for fila, rec in filas_motor.items():
        nit = norm_numeric_id(rec.get('NIT'))
        fac = norm_numeric_id(rec.get('FACTURA'))
        aut = norm_code(rec.get('AUTORIZACION'))
        if not nit or not fac:
            d['no_verificables'] += 1
            continue

        pos = fila - 2
        if 0 <= pos < len(src) and compatible(pos, nit, fac, aut):
            metodo = 'vinculo_fila'
        else:
            cands = [p for p in por_clave.get((nit, fac), []) if compatible(p, nit, fac, aut)]
            if not cands:
                d['sin_coincidencia'] += 1
                continue
            if len(cands) == 1:
                pos, metodo = cands[0], 'vinculo_clave'
            elif len({(s_usr[p], s_asi[p]) for p in cands}) == 1:
                pos, metodo = cands[0], 'vinculo_duplicado_identico'
            else:
                d['ambiguas'] += 1  # mismos NIT+factura con usuario/asiento distintos
                continue

        d[metodo] += 1
        meta[fila] = (s_usr[pos] or None, s_asi[pos] or None)

    d['filas_vinculadas'] = len(meta)
    d['filas_con_usuario'] = sum(1 for u, _ in meta.values() if u)
    d['filas_con_asiento'] = sum(1 for _, a in meta.values() if a)
    if d['ambiguas']:
        d['advertencias'].append(
            f"{d['ambiguas']} registro(s) SAP tienen NIT y factura repetidos con usuario/asiento distintos "
            'y no se pudo identificar cuál corresponde: quedaron vacíos.')
    if d['sin_coincidencia'] or d['no_verificables']:
        d['advertencias'].append(
            f"{d['sin_coincidencia'] + d['no_verificables']} registro(s) SAP no pudieron verificarse contra el "
            'archivo original (sin NIT/factura o sin coincidencia): usuario y asiento quedaron vacíos.')
    return meta, d



# ---------------------------------------------------------------------------
# Presentación de valores
# ---------------------------------------------------------------------------
def _show(value: Any) -> str:
    if _vacio(value):
        return '(vacío)'
    if isinstance(value, float):
        return f'{value:,.2f}'
    return str(value)


def _valores(item: dict) -> tuple[str, str, str]:
    """(VALOR EN SAP, VALOR EN SIAT, CLASIFICACIÓN POR CAMPO) alineados línea a línea."""
    sap, siat, clas = [], [], []
    for c in item['campos']:
        etq = f"{c['etiqueta']}"
        extra = ''
        if c['tag'] == 'REDONDEO':
            etq += ' [redondeo]'
        elif not _vacio(c.get('diferencia')):
            extra = f" (dif. {_num(c['diferencia']):,.2f})"
        sap.append(f"{etq}: {_show(c['valor_sap'])}")
        siat.append(f"{etq}: {_show(c['valor_siat'])}{extra}")
        clas.append(f"{c['etiqueta']}: {c['tag']}" + (f" · {c['motivo']}" if c['motivo'] else ''))
    return '\n'.join(sap), '\n'.join(siat), '\n'.join(clas)


# ---------------------------------------------------------------------------
# Escritura de hojas
# ---------------------------------------------------------------------------
def _put(ws, row: int, col: int, value: Any):
    """Escribe un valor; los textos nunca se interpretan como fórmula."""
    cell = ws.cell(row, col)
    if isinstance(value, float) and math.isnan(value):
        value = None
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub('', value)
        cell.value = value
        cell.data_type = 's'  # '=1+1' queda como texto, no como fórmula
        if value[:1] in ('=', '+', '-', '@', '\t', '\r'):
            cell.quotePrefix = True
    else:
        cell.value = value
    return cell


_BORDE = Border(bottom=Side(style='thin', color=BORDE))


def _bg(color: str) -> PatternFill:
    return PatternFill('solid', fgColor=color)


def _color_estado(estado: str) -> str:
    if estado.startswith('CRÍTICA'):
        return ROJO_SUAVE
    if estado == REV_IDENTIDAD:
        return MAGENTA_SUAVE
    if estado in ('REVISAR EMPAREJAMIENTO', 'SOLO SAP', 'SOLO SIAT'):
        return AZUL_SUAVE
    return AMARILLO_SUAVE  # diferencias menores (informativas)


def _tabla(ws, titulos, filas, anchos, color_pestana, dinero=(), texto=(), envolver=()):
    """Cabecera + datos + estilos funcionales (sin celdas combinadas). Columnas desde 0."""
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = color_pestana
    ws.row_dimensions[1].height = 34
    for i, (titulo, ancho) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(1, i, titulo)
        c.fill = _bg(MAGENTA if i == 2 else GRAFITO)
        c.font = Font(name=FUENTE, size=10, bold=True, color=BLANCO)
        c.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho

    f_base = Font(name=FUENTE, size=10, color=TEXTO)
    for r, fila in enumerate(filas, start=2):
        for j, valor in enumerate(fila):
            c = _put(ws, r, j + 1, valor)
            c.font = f_base
            c.border = _BORDE
            c.alignment = Alignment(vertical='top', horizontal='left', wrap_text=j in envolver)
            if j in dinero:
                c.number_format = NUM_FORMAT
                c.alignment = Alignment(vertical='top', horizontal='right')
            elif j in texto:
                c.number_format = '@'  # identificadores: nunca como número
        # Altura explícita según el texto envuelto (evita solapes y filas cortadas en Excel y LibreOffice).
        lineas = 1
        for j in envolver:
            if j < len(fila) and fila[j]:
                por_linea = max(8, int(anchos[j] * 0.82))
                n = sum(max(1, -(-len(t) // por_linea)) for t in str(fila[j]).split('\n'))
                lineas = max(lineas, n)
        ws.row_dimensions[r].height = min(409, max(16, 13 * lineas + 4))
        estado = ws.cell(r, 1)
        estado.fill = _bg(_color_estado(str(fila[0] or '')))
        estado.font = Font(name=FUENTE, size=10, bold=True, color=GRAFITO)
        b = ws.cell(r, 2)  # B: campos a corregir, en magenta
        b.font = Font(name=FUENTE, size=10, bold=True, color=MAGENTA)
        b.alignment = Alignment(vertical='top', horizontal='left', wrap_text=True)
        for col in (3, 4):  # C y D: trazabilidad SAP resaltada
            ws.cell(r, col).fill = _bg(AMARILLO_SUAVE)
            ws.cell(r, col).font = Font(name=FUENTE, size=10, bold=True, color=GRAFITO)

    ws.freeze_panes = 'E2'
    if filas:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(titulos))}{len(filas) + 1}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'


def _datos_factura(item: dict, auth: dict) -> dict:
    r = item['registro']
    nombre_sap, nombre_siat = r.get('RAZON SOCIAL SAP'), r.get('RAZON SOCIAL SIAT')
    return {
        'factura': r.get('FACTURA SAP'), 'nit': r.get('NIT SAP'),
        'razon_sap': nombre_sap,
        'razon_siat': nombre_siat if nombre_siat != nombre_sap else None,  # solo si difiere
        'fecha': r.get('FECHA SAP'), 'cuf': auth.get(item['fila_sap']),
        'imp_sap': r.get('IMPORTE SAP'), 'imp_siat': r.get('IMPORTE SIAT'),
        'dif': round(_num(r.get('IMPORTE SAP')) - _num(r.get('IMPORTE SIAT')), 2),
    }


def _hoja_control(wb, items: list[dict], meta: dict, auth: dict) -> int:
    titulos = [
        'ESTADO', 'CAMPOS CON DIFERENCIA', 'USUARIO SAP', 'N.º ASIENTO SAP',
        'N.º FACTURA', 'NIT', 'RAZÓN SOCIAL SAP', 'RAZÓN SOCIAL SIAT', 'FECHA',
        'CUF / AUTORIZACIÓN', 'IMPORTE SAP (Bs)', 'IMPORTE SIAT (Bs)', 'DIFERENCIA IMPORTE (Bs)',
        'VALOR EN SAP', 'VALOR EN SIAT', 'CLASIFICACIÓN POR CAMPO',
        'FILA SAP ORIGINAL', 'FILA SIAT ORIGINAL', 'MÉTODO EMPAREJAMIENTO',
    ]
    en_control = [i for i in items if i['estado']]
    en_control.sort(key=lambda i: (RANGO_ESTADO[i['estado']], -_num(i['registro'].get('IMPORTE SAP')),
                                   i['fila_sap'] or 0))
    filas = []
    for it in en_control:
        d = _datos_factura(it, auth)
        usuario, asiento = meta.get(it['fila_sap'], (None, None))
        v_sap, v_siat, clas = _valores(it)
        filas.append([
            it['estado'], it['registro'].get('CAMPOS CON DIFERENCIA') or '', usuario, asiento,
            d['factura'], d['nit'], d['razon_sap'], d['razon_siat'], d['fecha'], d['cuf'],
            d['imp_sap'], d['imp_siat'], d['dif'], v_sap, v_siat, clas,
            it['fila_sap'], it['fila_siat'], it['registro'].get('METODO EMPAREJAMIENTO'),
        ])
    ws = wb.create_sheet(HOJA_CONTROL)
    anchos = [27, 30, 17, 19, 14, 14, 36, 36, 12, 24, 16, 16, 16, 46, 46, 50, 12, 12, 26]
    _tabla(ws, titulos, filas, anchos, MAGENTA, dinero=(10, 11, 12), texto=(3, 4, 5, 9),
           envolver=(1, 6, 7, 13, 14, 15))

    # Resalta los datos concretos con diferencia: rojo = crítica, magenta suave = identidad,
    # amarillo = menor dentro de una factura que ya requiere revisión.
    columnas = {'FACTURA': (4,), 'NIT': (5,), 'RAZON_SOCIAL': (6, 7), 'FECHA': (8,),
                'AUTORIZACION': (9,), 'IMPORTE_TOTAL': (10, 11, 12)}
    color_tag = {TAG_CRITICA: ROJO_SUAVE, TAG_IDENTIDAD: MAGENTA_SUAVE, TAG_MENOR: AMARILLO_SUAVE}
    for r, it in enumerate(en_control, start=2):
        mayor = ROJO_SUAVE if it['estado'].startswith('CRÍTICA') else MAGENTA_SUAVE
        ws.cell(r, 14).fill = _bg(mayor)
        ws.cell(r, 15).fill = _bg(mayor)
        for c in it['campos']:
            if c['tag'] in color_tag and c['canon'] in columnas:
                for col in columnas[c['canon']]:
                    ws.cell(r, col + 1).fill = _bg(color_tag[c['tag']])
    return len(filas)


def _hoja_menores(wb, items: list[dict], meta: dict, auth: dict) -> int:
    titulos = [
        'TIPO DE DIFERENCIA', 'CAMPOS CON DIFERENCIA', 'USUARIO SAP', 'N.º ASIENTO SAP',
        'N.º FACTURA', 'NIT', 'RAZÓN SOCIAL SAP', 'RAZÓN SOCIAL SIAT', 'FECHA',
        'CUF / AUTORIZACIÓN', 'IMPORTE SAP (Bs)', 'IMPORTE SIAT (Bs)',
        'VALOR EN SAP', 'VALOR EN SIAT', 'MOTIVO DE CLASIFICACIÓN MENOR',
        'FILA SAP ORIGINAL', 'FILA SIAT ORIGINAL',
    ]
    menores = [i for i in items if i['nivel'] == 'MENOR']
    menores.sort(key=lambda i: (-_num(i['registro'].get('IMPORTE SAP')), i['fila_sap'] or 0))
    filas = []
    for it in menores:
        d = _datos_factura(it, auth)
        usuario, asiento = meta.get(it['fila_sap'], (None, None))
        v_sap, v_siat, _ = _valores(it)
        filas.append([
            tipo_menor(it), ' | '.join(str(c['etiqueta']) for c in it['campos']), usuario, asiento,
            d['factura'], d['nit'], d['razon_sap'], d['razon_siat'], d['fecha'], d['cuf'],
            d['imp_sap'], d['imp_siat'], v_sap, v_siat, '\n'.join(it['motivos_menor']),
            it['fila_sap'], it['fila_siat'],
        ])
    ws = wb.create_sheet(HOJA_MENORES)
    anchos = [24, 30, 17, 19, 14, 14, 36, 36, 12, 24, 16, 16, 44, 44, 56, 12, 12]
    _tabla(ws, titulos, filas, anchos, AMARILLO_PESTANA, dinero=(10, 11), texto=(3, 4, 5, 9),
           envolver=(1, 6, 7, 12, 13, 14))
    ws['A1'].comment = Comment(
        'Diferencias informativas de baja prioridad. Que una fila esté aquí NO significa que el dato '
        'esté corregido ni aprobado: solo que no requiere revisión prioritaria. Se separaron del '
        'CONTROL GENERAL por tratarse de la única diferencia de razón social o por solo redondeo. '
        'No supone aprobación de la identidad del proveedor.', 'Control SAP–SIAT')
    return len(filas)


def _hoja_pendientes(wb, solo_sap, solo_siat, candidatos, meta) -> int:
    """Cada hallazgo aparece UNA vez: un candidato absorbe a su fila solo-SAP y solo-SIAT."""
    titulos = [
        'SITUACION', 'MOTIVO / ACCION', 'USUARIO SAP', 'N.º ASIENTO SAP',
        'FACTURA SAP', 'FACTURA SIAT', 'NIT SAP', 'NIT SIAT',
        'PROVEEDOR SAP', 'PROVEEDOR SIAT', 'IMPORTE SAP (Bs)', 'IMPORTE SIAT (Bs)',
        'FECHA SAP', 'FECHA SIAT', 'PUNTAJE', 'FILA SAP ORIGINAL', 'FILA SIAT ORIGINAL',
        'EVIDENCIA / OBSERVACION',
    ]
    sap_en_candidato = {_key(c.get('FILA SAP')) for c in candidatos}
    siat_en_candidato = {_key(c.get('FILA SIAT CANDIDATA')) for c in candidatos}
    filas = []  # (grupo, importe, fila)
    for r in candidatos:
        f_sap = _key(r.get('FILA SAP'))
        usuario, asiento = meta.get(f_sap, (None, None))
        evidencia = f"Evidencia: {r.get('EVIDENCIA COINCIDENTE') or 'ninguna'}"
        if not _vacio(r.get('MARGEN VS 2DO')):
            evidencia += f" | margen vs 2.º candidato: {r.get('MARGEN VS 2DO')}"
        filas.append((0, _num(r.get('IMPORTE SAP')), [
            'REVISAR EMPAREJAMIENTO',
            'Posible cruce, NO aceptado: validar contra la factura física. Si no corresponde, ambas '
            'facturas quedan sin pareja.',
            usuario, asiento, r.get('FACTURA SAP'), r.get('FACTURA SIAT CANDIDATA'),
            r.get('NIT SAP'), r.get('NIT SIAT CANDIDATO'),
            r.get('PROVEEDOR SAP'), r.get('PROVEEDOR SIAT CANDIDATO'),
            r.get('IMPORTE SAP'), r.get('IMPORTE SIAT CANDIDATO'),
            r.get('FECHA SAP'), r.get('FECHA SIAT CANDIDATA'), r.get('PUNTAJE'),
            f_sap, _key(r.get('FILA SIAT CANDIDATA')), evidencia]))
    for r in solo_sap:
        f_sap = _key(r.get('_FILA_ORIGEN'))
        if f_sap in sap_en_candidato:
            continue
        usuario, asiento = meta.get(f_sap, (None, None))
        filas.append((1, _num(r.get('IMPORTE_TOTAL')), [
            'SOLO SAP', 'Sin correspondencia en SIAT', usuario, asiento,
            r.get('FACTURA'), None, r.get('NIT'), None, r.get('RAZON_SOCIAL'), None,
            r.get('IMPORTE_TOTAL'), None, r.get('FECHA'), None, None, f_sap, None,
            'Verificar el registro SAP y su documentación de respaldo']))
    for r in solo_siat:
        f_siat = _key(r.get('_FILA_ORIGEN'))
        if f_siat in siat_en_candidato:
            continue
        filas.append((2, _num(r.get('IMPORTE_TOTAL')), [
            'SOLO SIAT', 'Sin correspondencia en SAP', None, None,
            None, r.get('FACTURA'), None, r.get('NIT'), None, r.get('RAZON_SOCIAL'),
            None, r.get('IMPORTE_TOTAL'), None, r.get('FECHA'), None, None, f_siat,
            'Revisar si corresponde registrar la factura en SAP']))
    filas.sort(key=lambda x: (x[0], -x[1]))
    datos = [f[2] for f in filas]
    ws = wb.create_sheet(HOJA_PENDIENTES)
    anchos = [24, 40, 17, 19, 14, 14, 14, 14, 34, 34, 16, 16, 12, 12, 9, 12, 12, 56]
    _tabla(ws, titulos, datos, anchos, GRAFITO, dinero=(10, 11), texto=(3, 4, 5, 6, 7),
           envolver=(1, 17))
    return len(datos)


# ---------------------------------------------------------------------------
# Indicadores (para Streamlit; el Excel no lleva hoja de resumen)
# ---------------------------------------------------------------------------
def _indicadores(items, control, solo_sap, solo_siat, candidatos, diag) -> dict:
    n = Counter(i['nivel'] for i in items)
    estados = Counter(i['estado'] for i in items if i['estado'])
    menores = [i for i in items if i['nivel'] == 'MENOR']
    n_nominal = sum(1 for i in menores if any(c['tag'] == TAG_MENOR for c in i['campos']))
    n_red = sum(1 for i in menores if not any(c['tag'] == TAG_MENOR for c in i['campos']))
    return {
        'emparejadas': len(control),
        'sin_diferencias': n.get('OK', 0),
        'criticas': n.get('CRITICA', 0),
        'criticas_identificador': estados.get(CRIT_IDENT, 0),
        'revision_identidad': n.get('IDENTIDAD', 0),
        'menores_nominales': n_nominal,       # únicamente razón social, sin ninguna otra discrepancia
        'solo_redondeo': n_red,               # solo redondeo dentro de tolerancia
        'en_control_general': n.get('CRITICA', 0) + n.get('IDENTIDAD', 0),
        'en_menores': len(menores),
        'solo_sap': len(solo_sap),
        'solo_siat': len(solo_siat),
        'candidatos': len(candidatos),
        'total_sap': diag.get('total_sap', 0),
        'cobertura_usuario': diag.get('filas_con_usuario', 0),
        'cobertura_asiento': diag.get('filas_con_asiento', 0),
    }


# ---------------------------------------------------------------------------
# API pública
# ---------------------------------------------------------------------------
def generar_reporte_operativo(
    reporte_original: str | Path,
    sap_original: str | Path,
    reporte_salida: str | Path,
    columna_usuario: str | None = None,
    columna_asiento: str | None = None,
    tolerancia: float | None = None,
    equivalencias_path: str | Path | None = None,
    equivalencias_extra: Any = None,
) -> dict:
    """Genera el Excel compacto de 3 hojas y devuelve el diagnóstico (incluye ``indicadores``).

    ``reporte_original``: Excel producido por ``run_audit()`` (no se modifica).
    ``sap_original``: Excel SAP cargado por el usuario (fuente de usuario/asiento).
    ``columna_usuario`` / ``columna_asiento``: encabezado exacto para forzar la columna cuando la
    detección automática no reconoce la exportación.
    ``tolerancia``: tolerancia de redondeo usada en la auditoría (solo se cita en los motivos).
    ``equivalencias_path``: JSON de equivalencias de nombre verificadas (por NIT).
    ``equivalencias_extra``: equivalencias privadas adicionales (p. ej. secreto de Streamlit).
    """
    reporte_original, sap_original, reporte_salida = (
        Path(reporte_original), Path(sap_original), Path(reporte_salida)
    )
    # data_only=False: el Excel del motor no contiene fórmulas intencionales, pero escribe como
    # fórmula cualquier texto que empiece por '='; así se recupera el texto original (no un 0).
    raw = load_workbook(reporte_original, read_only=True, data_only=False)
    try:
        faltan = [s for s in HOJAS_MOTOR_REQUERIDAS if s not in raw.sheetnames]
        if faltan:
            raise ValueError('El motor no generó las hojas requeridas: ' + ', '.join(faltan))
        control = _read_sheet(raw, 'CONTROL GENERAL')
        diferencias = _read_sheet(raw, 'DETALLE DIFERENCIAS')
        redondeos = _read_sheet(raw, 'REDONDEOS')
        solo_sap = _read_sheet(raw, 'SOLO SAP')
        solo_siat = _read_sheet(raw, 'SOLO SIAT')
        candidatos = _read_sheet(raw, 'REVISAR EMPAREJAMIENTO')
        base_sap = _read_sheet(raw, 'BASE SAP NORMALIZADA')
        reparaciones = _read_sheet(raw, 'REPARACIONES SIAT')
    finally:
        raw.close()

    meta, diag = _vincular_usuario_asiento(sap_original, base_sap, columna_usuario, columna_asiento)
    auth = {_key(r.get('_FILA_ORIGEN')): r.get('AUTORIZACION') for r in base_sap}
    items = clasificar_facturas(control, diferencias, redondeos,
                                cargar_equivalencias(equivalencias_path, equivalencias_extra), tolerancia)
    if reparaciones:
        lineas = ', '.join(str(r.get('LINEAS FISICAS')) for r in reparaciones[:15])
        mas = f' y {len(reparaciones) - 15} más' if len(reparaciones) > 15 else ''
        diag['advertencias'].append(
            f'El CSV SIAT tuvo {len(reparaciones)} incidencia(s) estructural(es) reparada(s) '
            f'automáticamente por el motor (líneas físicas: {lineas}{mas}). Verifique esos registros.')

    out = Workbook()
    out.remove(out.active)
    n_control = _hoja_control(out, items, meta, auth)
    n_pend = _hoja_pendientes(out, solo_sap, solo_siat, candidatos, meta)
    n_menores = _hoja_menores(out, items, meta, auth)

    # Control de integridad: ninguna factura emparejada ni sin pareja se pierde en la capa de reporte.
    ind = _indicadores(items, control, solo_sap, solo_siat, candidatos, diag)
    suma = (ind['sin_diferencias'] + ind['criticas'] + ind['revision_identidad']
            + ind['menores_nominales'] + ind['solo_redondeo'])
    esperado_pend = len(candidatos) + len(solo_sap) + len(solo_siat) - (
        len({_key(c.get('FILA SAP')) for c in candidatos} & {_key(r.get('_FILA_ORIGEN')) for r in solo_sap})
        + len({_key(c.get('FILA SIAT CANDIDATA')) for c in candidatos}
              & {_key(r.get('_FILA_ORIGEN')) for r in solo_siat}))
    if (suma != len(control) or n_control != ind['en_control_general'] or n_menores != ind['en_menores']
            or n_pend != esperado_pend):
        raise RuntimeError('El reporte compacto no conserva todos los registros del motor.')

    out.properties.creator = __author__
    out.properties.title = 'Control Integral de Compras SAP–SIAT (reporte operativo)'
    out.properties.description = f'Generado por {__project__} v{__version__}'
    reporte_salida.parent.mkdir(parents=True, exist_ok=True)
    out.save(reporte_salida)
    out.close()

    diag['indicadores'] = ind
    diag['nombre_archivo'] = reporte_salida.name
    diag['total_conciliadas'] = len(control)
    diag['total_pendientes'] = n_pend
    return diag
