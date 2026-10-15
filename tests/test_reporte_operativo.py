"""Pruebas del reporte operativo (capa posterior a run_audit) con datos SINTÉTICOS."""
from __future__ import annotations

import hashlib
import subprocess
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

import reporte_operativo as ro
from motor_auditoria import run_audit
from sinteticos import escenario, escribir_sap, sap_row

RAIZ = Path(__file__).resolve().parents[1]
COLS_BASE = list(sap_row('1', 1, 'A', '2025-01-01', 1, 'x'))


def _generar(tmp_path, **kw):
    sap, siat = escenario(tmp_path, **kw)
    motor_xlsx = tmp_path / 'motor.xlsx'
    _, metrics = run_audit(sap, siat, output_path=motor_xlsx)
    salida = tmp_path / 'reporte.xlsx'
    diag = ro.generar_reporte_operativo(motor_xlsx, sap, salida)
    return openpyxl.load_workbook(salida), metrics, diag, motor_xlsx, sap


@pytest.fixture
def corrida(tmp_path):
    return _generar(tmp_path)


def _filas(ws):
    cab = [c.value for c in ws[1]]
    return [dict(zip(cab, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


# ---------------------------------------------------------------- estructura
HOJAS = ['CONTROL GENERAL', 'POR CONCILIAR', 'DIFERENCIAS MENORES']


def test_exactamente_tres_hojas_en_orden(corrida):
    wb = corrida[0]
    assert wb.sheetnames == HOJAS
    assert 'RESUMEN' not in wb.sheetnames


def test_columnas_A_a_D_control_general(corrida):
    ws = corrida[0]['CONTROL GENERAL']
    assert [ws.cell(1, i).value for i in range(1, 5)] == [
        'ESTADO', 'CAMPOS CON DIFERENCIA', 'USUARIO SAP', 'N.º ASIENTO SAP']
    f = {r['N.º FACTURA']: r for r in _filas(ws)}
    assert f['102']['ESTADO'] == 'CRÍTICA · IMPORTE O TRIBUTARIO'
    assert 'Importe total compra' in f['102']['CAMPOS CON DIFERENCIA']
    assert (f['102']['USUARIO SAP'], f['102']['N.º ASIENTO SAP']) == ('LGOMEZ', '1900000102')
    assert f['107']['ESTADO'] == 'CRÍTICA · IDENTIFICADOR'
    assert f['104']['ESTADO'] == 'CRÍTICA · FECHA'


def test_control_general_solo_diferencias_relevantes(corrida):
    wb = corrida[0]
    facturas = {r['N.º FACTURA'] for r in _filas(wb['CONTROL GENERAL'])}
    assert facturas == {'102', '104', '107'}  # no OK (101), no menores (103)
    assert {r['ESTADO'] for r in _filas(wb['CONTROL GENERAL'])} <= set(
        ['CRÍTICA · IDENTIFICADOR', 'CRÍTICA · IMPORTE O TRIBUTARIO', 'CRÍTICA · FECHA', 'REVISIÓN DE IDENTIDAD'])
    menores = {r['N.º FACTURA']: r for r in _filas(wb['DIFERENCIAS MENORES'])}
    assert set(menores) == {'103'} and menores['103']['TIPO DE DIFERENCIA'] == 'REDONDEO'
    assert 'dentro de la tolerancia' in menores['103']['MOTIVO DE CLASIFICACIÓN MENOR']


def test_columna_B_es_la_de_la_hoja_original_del_motor(corrida):
    wb, _, _, motor_xlsx, _ = corrida
    original = {str(r['FACTURA SAP']): r['CAMPOS CON DIFERENCIA'] or ''
                for r in _filas(openpyxl.load_workbook(motor_xlsx)['CONTROL GENERAL'])}
    nuevo = {str(r['N.º FACTURA']): r['CAMPOS CON DIFERENCIA'] or '' for r in _filas(wb['CONTROL GENERAL'])}
    assert nuevo and all(original[k] == v for k, v in nuevo.items())


def test_ceros_iniciales_del_asiento_se_conservan_como_texto(corrida):
    ws = corrida[0]['CONTROL GENERAL']
    f = {r['N.º FACTURA']: r for r in _filas(ws)}
    assert f['104']['N.º ASIENTO SAP'] == '0001900104'
    celda = next(c for c in ws['D'][1:] if c.value == '0001900104')
    assert celda.data_type == 's' and celda.number_format == '@'


def test_valores_sap_y_siat_por_campo(corrida):
    f = {r['N.º FACTURA']: r for r in _filas(corrida[0]['CONTROL GENERAL'])}
    r = f['102']
    assert 'Importe total compra: 250' in r['VALOR EN SAP'].replace('.00', '') or 'Importe total compra: 250.00' in r['VALOR EN SAP']
    assert 'Importe total compra: 205' in r['VALOR EN SIAT'].replace('.00', '')
    assert 'Razón social: PROVEEDOR DOS SA' in r['VALOR EN SAP'] and 'S.A.' in r['VALOR EN SIAT']
    assert 'Razón social: MENOR' in r['CLASIFICACIÓN POR CAMPO']       # menor visible, no oculta lo crítico
    assert 'Importe total compra: CRÍTICA' in r['CLASIFICACIÓN POR CAMPO']
    assert r['DIFERENCIA IMPORTE (Bs)'] == 45
    assert r['CUF / AUTORIZACIÓN'] == 'AUT0000102'
    assert r['RAZÓN SOCIAL SIAT'] == 'PROVEEDOR DOS S.A.'


# ----------------------------------------------------------- integridad datos
def test_no_se_pierde_ninguna_factura(corrida):
    wb, m, diag, motor_xlsx, _ = corrida
    control = _filas(wb['CONTROL GENERAL'])
    menores = _filas(wb['DIFERENCIAS MENORES'])
    pend = _filas(wb['POR CONCILIAR'])
    ind = diag['indicadores']
    # Todo lo emparejado está en exactamente un lugar (control, menores) o es OK (no se exporta)
    assert len(control) + len(menores) + ind['sin_diferencias'] == m['matched']
    assert not {r['FILA SAP ORIGINAL'] for r in control} & {r['FILA SAP ORIGINAL'] for r in menores}
    # Todo lo no emparejado del motor aparece UNA sola vez en POR CONCILIAR
    mw = openpyxl.load_workbook(motor_xlsx)
    sap_sin = {r['_FILA_ORIGEN'] for r in _filas(mw['SOLO SAP'])}
    siat_sin = {r['_FILA_ORIGEN'] for r in _filas(mw['SOLO SIAT'])}
    sap_vistas = [r['FILA SAP ORIGINAL'] for r in pend if r['FILA SAP ORIGINAL'] is not None]
    siat_vistas = [r['FILA SIAT ORIGINAL'] for r in pend if r['FILA SIAT ORIGINAL'] is not None]
    assert set(sap_vistas) == sap_sin and set(siat_vistas) >= {s for s in siat_sin}
    assert len(sap_vistas) == len(set(sap_vistas))                      # sin duplicar hallazgos
    # Todas las filas SAP del motor siguen rastreables (control + menores + OK + sin pareja)
    assert len(control) + len(menores) + ind['sin_diferencias'] + len(sap_sin) == m['sap']
    assert diag['total_conciliadas'] == m['matched']


def test_indicadores_suman_y_no_alteran_las_cifras_del_motor(corrida):
    wb, m, diag, *_ = corrida
    ind = diag['indicadores']
    assert ind['emparejadas'] == m['matched'] and ind['sin_diferencias'] == m['ok']
    assert (ind['sin_diferencias'] + ind['criticas'] + ind['revision_identidad']
            + ind['menores_nominales'] + ind['solo_redondeo']) == m['matched']
    assert ind['solo_sap'] == m['solo_sap'] and ind['solo_siat'] == m['solo_siat']
    assert ind['candidatos'] == m['review']
    assert ind['en_control_general'] == len(_filas(wb['CONTROL GENERAL']))
    assert ind['en_menores'] == len(_filas(wb['DIFERENCIAS MENORES']))
    # lo que el motor llama diferencias/redondeos se reparte sin perder facturas
    assert ind['criticas'] + ind['revision_identidad'] + ind['menores_nominales'] + ind['solo_redondeo'] == m['diff'] + m['rounding']


def test_importes_conservan_valores_originales(corrida):
    wb, _, _, motor_xlsx, _ = corrida
    orig = {(r['FACTURA SAP'], r['FILA SAP']): (r['IMPORTE SAP'], r['IMPORTE SIAT'])
            for r in _filas(openpyxl.load_workbook(motor_xlsx)['CONTROL GENERAL'])}
    for hoja in ('CONTROL GENERAL', 'DIFERENCIAS MENORES'):
        for r in _filas(wb[hoja]):
            assert orig[(r['N.º FACTURA'], r['FILA SAP ORIGINAL'])] == (r['IMPORTE SAP (Bs)'], r['IMPORTE SIAT (Bs)'])
    menor = _filas(wb['DIFERENCIAS MENORES'])[0]
    assert (menor['IMPORTE SAP (Bs)'], menor['IMPORTE SIAT (Bs)']) == (500, 500.03)  # sin redondear
    assert wb['CONTROL GENERAL']['K2'].number_format.startswith('#,##0.00')


def test_motor_y_su_excel_permanecen_intactos(tmp_path):
    sap, siat = escenario(tmp_path)
    motor_xlsx = tmp_path / 'motor.xlsx'
    run_audit(sap, siat, output_path=motor_xlsx)
    antes = hashlib.sha256(motor_xlsx.read_bytes()).hexdigest()
    ro.generar_reporte_operativo(motor_xlsx, sap, tmp_path / 'r.xlsx')
    assert hashlib.sha256(motor_xlsx.read_bytes()).hexdigest() == antes
    # el motor no fue tocado respecto del commit de partida de esta fase
    if subprocess.run(['git', 'cat-file', '-e', 'd0c07e1'], cwd=RAIZ, capture_output=True).returncode != 0:
        pytest.skip('commit de referencia no disponible en este checkout')
    for archivo in ('motor_auditoria.py',):
        assert subprocess.run(['git', 'diff', '--quiet', 'd0c07e1', '--', archivo], cwd=RAIZ).returncode == 0


# --------------------------------------------------------------- ordenamiento
def test_orden_por_gravedad_y_luego_importe_descendente(corrida):
    filas = _filas(corrida[0]['CONTROL GENERAL'])
    rango = {'CRÍTICA · IDENTIFICADOR': 0, 'CRÍTICA · IMPORTE O TRIBUTARIO': 1,
             'CRÍTICA · FECHA': 2, 'REVISIÓN DE IDENTIDAD': 3}
    claves = [(rango[r['ESTADO']], -r['IMPORTE SAP (Bs)']) for r in filas]
    assert claves == sorted(claves)
    assert [r['N.º FACTURA'] for r in filas] == ['107', '102', '104']


def test_filtros_paneles_y_formato(corrida):
    wb = corrida[0]
    for nombre in HOJAS:
        ws = wb[nombre]
        assert ws.freeze_panes == 'E2'  # A:D siempre visibles
        assert ws.auto_filter.ref == f'A1:{ws.cell(1, ws.max_column).column_letter}{ws.max_row}'
        assert not ws.merged_cells.ranges                                  # sin celdas combinadas
    ws = wb['CONTROL GENERAL']
    assert ws['B1'].fill.fgColor.rgb.endswith('C5004E')  # magenta institucional
    assert ws['A1'].fill.fgColor.rgb.endswith('31343A')  # grafito
    assert ws.sheet_properties.tabColor.rgb.endswith('C5004E')
    men = wb['DIFERENCIAS MENORES']
    assert men.sheet_properties.tabColor.rgb.endswith('FFE699')          # amarillo suave, última pestaña
    assert men['A2'].fill.fgColor.rgb.endswith(ro.AMARILLO_SUAVE)
    assert 'NO significa' in men['A1'].comment.text
    assert ws['A2'].fill.fgColor.rgb.endswith(ro.ROJO_SUAVE)


def test_por_conciliar_agrupa_y_no_acepta_candidatos(corrida):
    pend = _filas(corrida[0]['POR CONCILIAR'])
    orden = [r['SITUACION'] for r in pend]
    assert orden == sorted(orden, key={'REVISAR EMPAREJAMIENTO': 0, 'SOLO SAP': 1, 'SOLO SIAT': 2}.get)
    for r in pend:
        if r['SITUACION'] == 'REVISAR EMPAREJAMIENTO':
            assert 'NO aceptado' in r['MOTIVO / ACCION']
        if r['SITUACION'] == 'SOLO SIAT':
            assert r['USUARIO SAP'] is None and r['N.º ASIENTO SAP'] is None
    # un candidato absorbe su fila solo-SAP/solo-SIAT: ningún hallazgo aparece dos veces
    candidatos_sap = {r['FILA SAP ORIGINAL'] for r in pend if r['SITUACION'] == 'REVISAR EMPAREJAMIENTO'}
    assert not candidatos_sap & {r['FILA SAP ORIGINAL'] for r in pend if r['SITUACION'] == 'SOLO SAP'}
    en_control = {r['FILA SAP ORIGINAL'] for r in _filas(corrida[0]['CONTROL GENERAL'])}
    assert not en_control & {r['FILA SAP ORIGINAL'] for r in pend}


# ------------------------------------------------- usuario / asiento: seguridad
def test_duplicados_se_resuelven_por_fila_y_no_se_mezclan(corrida):
    """Dos filas SAP con mismo NIT+factura y asientos distintos: cada una conserva el suyo."""
    pend = _filas(corrida[0]['POR CONCILIAR'])
    dup = {r['FILA SAP ORIGINAL']: (r['USUARIO SAP'], r['N.º ASIENTO SAP'])
           for r in pend if r['FACTURA SAP'] == '108'}
    assert dup == {9: ('USRA', '1900000108'), 10: ('USRB', '1900000109')}
    assert corrida[2]['ambiguas'] == 0 and corrida[2]['advertencias'] == []


def _base(fila, nit, fac, aut=''):
    return {'_FILA_ORIGEN': fila, 'NIT': nit, 'FACTURA': fac, 'AUTORIZACION': aut}


def _fuente(tmp_path, filas, nombre='src.xlsx'):
    return escribir_sap(tmp_path / nombre, filas)


def test_ambiguo_sin_pista_de_fila_queda_vacio(tmp_path):
    """Si la fila de origen no verifica y la clave está duplicada con datos distintos: vacío."""
    src = _fuente(tmp_path, [
        sap_row('1000000008', 108, 'A8', '2025-01-16', 60, 'X', 'USRA', '111'),
        sap_row('1000000008', 108, 'A8', '2025-01-16', 60, 'X', 'USRB', '222'),
    ])
    meta, d = ro._vincular_usuario_asiento(src, [_base(99, '1000000008', '108', 'A8')])
    assert meta == {} and d['ambiguas'] == 1
    assert any('NIT y factura repetidos' in a for a in d['advertencias'])


def test_duplicado_identico_se_asigna(tmp_path):
    src = _fuente(tmp_path, [
        sap_row('1000000008', 108, 'A8', '2025-01-16', 60, 'X', 'USRA', '111'),
        sap_row('1000000008', 108, 'A8', '2025-01-16', 60, 'X', 'USRA', '111'),
    ])
    meta, d = ro._vincular_usuario_asiento(src, [_base(99, '1000000008', '108', 'A8')])
    assert meta == {99: ('USRA', '111')} and d['vinculo_duplicado_identico'] == 1


def test_clave_unica_con_desfase_de_fila(tmp_path):
    src = _fuente(tmp_path, [sap_row('1000000008', 108, 'A8', '2025-01-16', 60, 'X', 'USRA', '111')])
    meta, d = ro._vincular_usuario_asiento(src, [_base(7, '1000000008', '108')])
    assert meta == {7: ('USRA', '111')} and d['vinculo_clave'] == 1


def test_fila_que_no_verifica_nit_o_factura_no_se_asigna(tmp_path):
    src = _fuente(tmp_path, [sap_row('1000000001', 101, 'A1', '2025-01-10', 10, 'X', 'USRA', '111')])
    for nit, fac in (('1000000002', '101'), ('1000000001', '999')):
        meta, d = ro._vincular_usuario_asiento(src, [_base(2, nit, fac)])
        assert meta == {} and d['sin_coincidencia'] == 1


def test_autorizacion_contradictoria_impide_vinculo(tmp_path):
    src = _fuente(tmp_path, [sap_row('1000000001', 101, 'AUT-A', '2025-01-10', 10, 'X', 'USRA', '111')])
    meta, _ = ro._vincular_usuario_asiento(src, [_base(2, '1000000001', '101', 'AUT-B')])
    assert meta == {}
    meta, _ = ro._vincular_usuario_asiento(src, [_base(2, '1000000001', '101', 'AUT-A')])
    assert meta == {2: ('USRA', '111')}


def test_nit_con_ceros_y_formatos_numericos(tmp_path):
    src = _fuente(tmp_path, [sap_row('00123456', 104, 'A', '2025-01-10', 10, 'X', 'U', 77)])
    meta, _ = ro._vincular_usuario_asiento(src, [_base(2, '123456', '104')])
    assert meta == {2: ('U', '77')}


def test_sin_columnas_usuario_asiento_deja_vacio_y_advierte(tmp_path):
    cols = [c for c in COLS_BASE if c not in ('Usuario', 'N° Asiento')]
    wb, m, diag, *_ = _generar(tmp_path, columnas=cols)
    f = _filas(wb['CONTROL GENERAL']) + _filas(wb['DIFERENCIAS MENORES'])
    assert f and all(r['USUARIO SAP'] is None and r['N.º ASIENTO SAP'] is None for r in f)
    assert len(f) + diag['indicadores']['sin_diferencias'] == m['matched']  # el reporte sigue completo
    assert any('usuario' in a.lower() for a in diag['advertencias'])
    assert any('Encabezados de la hoja SAP leída' in a for a in diag['advertencias'])
    assert diag['filas_con_usuario'] == 0 and diag['filas_con_asiento'] == 0


def test_solo_asiento_o_solo_usuario(tmp_path):
    wb, _, diag, *_ = _generar(tmp_path, columnas=[c for c in COLS_BASE if c != 'Usuario'])
    f = {r['N.º FACTURA']: r for r in _filas(wb['CONTROL GENERAL'])}
    assert f['102']['USUARIO SAP'] is None and f['102']['N.º ASIENTO SAP'] == '1900000102'
    assert diag['asiento_detectado'] and not diag['usuario_detectado']


@pytest.mark.parametrize('u,a', [
    ('Usuario Registro', 'Nº Asiento'), ('Creado por', 'Doc. Contable'),
    ('ERNAM', 'BELNR'), ('USUARIO', 'N.º Asiento Contable'),
    ('Usuario SAP', 'Documento Contable'),
])
def test_alias_de_encabezados(tmp_path, u, a):
    sap, siat = escenario(tmp_path)
    df = pd.read_excel(sap, dtype=object).rename(columns={'Usuario': u, 'N° Asiento': a})
    df.to_excel(sap, index=False)
    motor_xlsx = tmp_path / 'm.xlsx'
    run_audit(sap, siat, output_path=motor_xlsx)
    diag = ro.generar_reporte_operativo(motor_xlsx, sap, tmp_path / 'r.xlsx')
    assert (diag['columna_usuario'], diag['columna_asiento']) == (u, a)
    assert diag['filas_con_usuario'] == 8 and diag['filas_con_asiento'] == 8


@pytest.mark.parametrize('encabezado,esperado', [
    ('Usuario', 'usuario'), ('Usuario Registro', 'usuario'), ('Creado por', 'usuario'),
    ('USNAM', 'usuario'), ('N° Asiento', 'asiento'), ('Nº de asiento', 'asiento'),
    ('Documento Contable', 'asiento'), ('BELNR', 'asiento'),
    # NO deben confundirse:
    ('Nro. Factura', None), ('Fec. Factura', None), ('NIT Proveedor', None),
    ('Precio Unitario', None), ('Nro. DUI/DIM', None), ('Cod. Autorización', None),
    ('Doc. Factura', None), ('Usuario Modificación', None), ('Aprobado por usuario', None),
    ('Asiento Factura', None), ('Cód. Control', None), ('Doc. Anulacion', None),
    ('Doc. Logístico', None), ('Cl. Documento', None), ('Fec. Modificación', None),
    ('Cod. Acreedor', None), ('Nro. Documento Factura', None),
    # Encabezados de la exportación SAP real (sep-2026):
    ('Nro. Documento', 'asiento'),
])
def test_clasificacion_de_encabezados(encabezado, esperado):
    assert ro._classify_header(encabezado) == esperado


def test_varias_columnas_de_usuario_ambiguas_no_se_eligen():
    sel, av = ro.detectar_columnas_sap(['Usuario Alta', 'Usuario Baja', 'N° Asiento'])
    assert sel == {'usuario': None, 'asiento': 'N° Asiento'} and av
    sel, _ = ro.detectar_columnas_sap(['Usuario', 'Creado por'])  # preferida por semántica
    assert sel['usuario'] == 'Creado por'


def test_asiento_igual_a_factura_se_descarta(tmp_path):
    filas = [sap_row('1000000001', 100 + i, f'A{i}', '2025-01-10', 10 + i, 'X', 'U', str(100 + i))
             for i in range(5)]
    src = _fuente(tmp_path, filas)
    meta, d = ro._vincular_usuario_asiento(
        src, [_base(i + 2, '1000000001', str(100 + i), f'A{i}') for i in range(5)])
    assert all(a is None for _, a in meta.values())
    assert any('número de factura' in a for a in d['advertencias'])


def test_override_manual_de_columnas(tmp_path):
    sap, siat = escenario(tmp_path)
    df = pd.read_excel(sap, dtype=object).rename(columns={'Usuario': 'Operador', 'N° Asiento': 'Referencia'})
    df.to_excel(sap, index=False)
    m = tmp_path / 'm.xlsx'
    run_audit(sap, siat, output_path=m)
    sin = ro.generar_reporte_operativo(m, sap, tmp_path / 'a.xlsx')
    con = ro.generar_reporte_operativo(m, sap, tmp_path / 'b.xlsx', 'Operador', 'Referencia')
    assert sin['filas_con_usuario'] == 0 and con['filas_con_usuario'] == 8 and con['filas_con_asiento'] == 8


# ------------------------------------------------------ inyección de fórmulas
def test_formulas_externas_no_se_interpretan(tmp_path):
    peligro = sap_row('1000000009', 109, 'AUT0000109', '2025-01-17', 70.0, '=HYPERLINK("http://x","clic")',
                      usuario='=1+1', asiento='@SUM(A1)')
    wb, *_ = _generar(tmp_path, extra_sap=[peligro])
    for ws in wb:
        for fila in ws.iter_rows():
            for c in fila:
                assert c.data_type != 'f', f'{ws.title}!{c.coordinate} es fórmula'
    pend = {r['FACTURA SAP']: r for r in _filas(wb['POR CONCILIAR']) if r['SITUACION'] == 'SOLO SAP'}
    r = pend['109']
    assert r['USUARIO SAP'] == '=1+1' and r['N.º ASIENTO SAP'] == '@SUM(A1)'
    assert r['PROVEEDOR SAP'].startswith('=HYPERLINK')
    celda = next(c for c in wb['POR CONCILIAR']['C'] if c.value == '=1+1')
    assert celda.data_type == 's' and celda.quotePrefix


# -------------------------------------------------------- casos de borde / app
def test_sin_ninguna_conciliacion_genera_hojas_vacias_validas(tmp_path):
    sap = escribir_sap(tmp_path / 's.xlsx', [sap_row('1000000001', 1, 'AA', '2025-01-01', 10, 'X')])
    from sinteticos import escribir_siat, siat_row
    siat = escribir_siat(tmp_path / 't.csv', [siat_row(1, '9999999999', 77, 'ZZ000000', '2025-02-02', 99, 'Y')])
    m = tmp_path / 'm.xlsx'
    run_audit(sap, siat, output_path=m)
    ro.generar_reporte_operativo(m, sap, tmp_path / 'r.xlsx')
    wb = openpyxl.load_workbook(tmp_path / 'r.xlsx')
    assert wb.sheetnames == HOJAS
    assert wb['CONTROL GENERAL'].max_row == 1 and wb['DIFERENCIAS MENORES'].max_row == 1
    assert wb['POR CONCILIAR'].max_row == 3


def test_motor_incompleto_se_rechaza(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = 'OTRA'
    wb.save(tmp_path / 'x.xlsx')
    with pytest.raises(ValueError, match='hojas requeridas'):
        ro.generar_reporte_operativo(tmp_path / 'x.xlsx', tmp_path / 'x.xlsx', tmp_path / 'y.xlsx')


def test_streamlit_app_ejecuta_sin_errores(tmp_path):
    pytest.importorskip('streamlit')
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(RAIZ / 'app.py'), default_timeout=60).run()
    assert not at.exception
    assert any('Carga el Excel de SAP y el CSV del SIAT' in i.value for i in at.info)


def test_encabezados_reales_de_la_exportacion_sap():
    """Los 32 encabezados reales (solo nombres) deben dar exactamente Usuario y Nro. Documento."""
    reales = ['N°', 'Especificación', 'NIT Proveedor', 'Razón Social Proveedor', 'Cod. Autorización',
              'Nro. Factura', 'Nro. DUI/DIM', 'Fec. Factura', 'Importe Total Compra', 'Importe ICE',
              'Importe IEHD', 'Importe IPJ', 'Tasas', 'Otro no Sujeto a CF', 'Importes Exentos',
              'Importe C.G. Tasa Cero', 'Sub total', 'Desc. Bon. y Rebajas Obtenidas', 'Importe Gift Card',
              'Importe Base CF', 'Crédito Fiscal IVA', 'Tipo de Compra', 'Cód. Control', 'Nro. Documento',
              'Cl. Documento', 'Fec. Contabilización', 'Fec. Modificación', 'Doc. Anulacion', 'Usuario',
              'Cod. Acreedor', 'Texto Posición', 'Doc. Logístico']
    sel, av = ro.detectar_columnas_sap(reales)
    assert sel == {'usuario': 'Usuario', 'asiento': 'Nro. Documento'} and av == []


# ------------------------------------------------------------- formatos .xls / .xlsx
def _resumen_hojas(ruta):
    wb = openpyxl.load_workbook(ruta)
    return {ws.title: [[c.value for c in f] for f in ws.iter_rows()] for ws in wb}


def _corrida_con_sap(tmp_path, sap_alterno):
    _, siat = escenario(tmp_path)
    motor = tmp_path / 'm_alt.xlsx'
    _, m = run_audit(sap_alterno, siat, output_path=motor)
    diag = ro.generar_reporte_operativo(motor, sap_alterno, tmp_path / 'r_alt.xlsx')
    return m, diag, tmp_path / 'r_alt.xlsx'


def test_extension_incorrecta_xlsx_guardado_como_xls(tmp_path):
    """Un .xlsx renombrado a .xls (o al revés) se lee por contenido, igual que el motor."""
    sap, _ = escenario(tmp_path)
    base_m, base_d, base_r = _corrida_con_sap(tmp_path, sap)
    falso = tmp_path / 'SAP_falso.xls'
    falso.write_bytes(sap.read_bytes())
    m, d, r = _corrida_con_sap(tmp_path, falso)
    assert m == base_m and d['filas_con_usuario'] == base_d['filas_con_usuario'] == 8
    assert _resumen_hojas(r) == _resumen_hojas(base_r)


def test_xls_binario_real_da_el_mismo_reporte_que_xlsx(tmp_path):
    """.xls BIFF real (convertido con LibreOffice): mismo resultado, usuario/asiento y ceros incluidos."""
    import shutil
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        pytest.skip('LibreOffice no disponible para generar un .xls de prueba')
    pytest.importorskip('xlrd')
    sap, _ = escenario(tmp_path)
    base_m, base_d, base_r = _corrida_con_sap(tmp_path, sap)
    out = tmp_path / 'xls'
    out.mkdir()
    subprocess.run([soffice, '--headless', '--convert-to', 'xls:MS Excel 97', '--outdir', str(out), str(sap)],
                   check=True, capture_output=True, timeout=180)
    xls = out / 'SAP.xls'
    assert xls.read_bytes()[:4] == b'\xd0\xcf\x11\xe0'  # contenedor OLE2 = .xls real
    m, d, r = _corrida_con_sap(tmp_path, xls)
    assert m == base_m and d['filas_con_usuario'] == 8 and d['filas_con_asiento'] == 8
    wb = openpyxl.load_workbook(r)
    f = {x['N.º FACTURA']: x for x in _filas(wb['CONTROL GENERAL'])}
    assert f['104']['N.º ASIENTO SAP'] == '0001900104'  # ceros iniciales intactos
    assert _resumen_hojas(r) == _resumen_hojas(base_r)
