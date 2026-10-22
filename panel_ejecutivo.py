"""Centro ejecutivo Streamlit, solo lectura, sobre el Excel operativo SAP-SIAT.

No modifica el motor, conciliaciones ni datos de entrada.
"""
from collections import Counter
from html import escape
from io import BytesIO

import streamlit as st
from openpyxl import load_workbook


def _filas(wb, nombre):
    rows = wb[nombre].iter_rows(values_only=True)
    cab = next(rows)
    return [dict(zip(cab, r)) for r in rows if any(v is not None for v in r)]


def _texto(v):
    return escape(str(v)) if v is not None and str(v) != "" else "—"


def _cantidad(v):
    return f"{int(v):,}".replace(",", ".")


def _bolivianos(v):
    try:
        return f"Bs {float(v):,.2f}".replace(",", "#").replace(".", ",").replace("#", ".")
    except (ValueError, TypeError):
        return "—"


def _numero(r, tipo):
    return (r.get("FACTURA SAP") or r.get("FACTURA SIAT")) if tipo == "pend" else r.get("N.º FACTURA")


def _importe(r, tipo):
    val = r.get("IMPORTE SAP (Bs)")
    if val is None and tipo == "pend":
        val = r.get("IMPORTE SIAT (Bs)")
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def _campos(texto):
    resultado = {}
    for linea in str(texto or "").splitlines():
        if ":" in linea:
            campo, valor = linea.split(":", 1)
            resultado[campo.strip()] = valor.strip()
    return resultado


def _detalle(r, tipo):
    with st.container(border=True):
        asiento, usuario, fecha = st.columns([1, 1, 1])
        asiento.caption("ASIENTO SAP")
        asiento.markdown("**" + _texto(r.get("N.º ASIENTO SAP")) + "**", unsafe_allow_html=True)
        usuario.caption("USUARIO SAP")
        usuario.markdown("**" + _texto(r.get("USUARIO SAP")) + "**", unsafe_allow_html=True)
        fecha.caption("FECHA DOCUMENTO")
        fecha.markdown("**" + _texto(r.get("FECHA") or r.get("FECHA SAP") or r.get("FECHA SIAT")) + "**", unsafe_allow_html=True)

        if tipo == "pend":
            sap, siat = st.columns(2)
            with sap:
                st.markdown("**REGISTRO SAP**")
                st.caption("Factura " + _texto(r.get("FACTURA SAP")))
                st.write(r.get("PROVEEDOR SAP") or "Sin correspondencia")
                st.markdown("**" + _bolivianos(r.get("IMPORTE SAP (Bs)")) + "**")
            with siat:
                st.markdown("**REGISTRO SIAT**")
                st.caption("Factura " + _texto(r.get("FACTURA SIAT")))
                st.write(r.get("PROVEEDOR SIAT") or "Sin correspondencia")
                st.markdown("**" + _bolivianos(r.get("IMPORTE SIAT (Bs)")) + "**")
            st.info(str(r.get("MOTIVO / ACCION") or "Requiere validación manual."))
            return

        st.caption("CUF / AUTORIZACIÓN")
        st.markdown("<div class='ue-cuf'>" + _texto(r.get("CUF / AUTORIZACIÓN")) + "</div>", unsafe_allow_html=True)
        nombres = [x.strip() for x in str(r.get("CAMPOS CON DIFERENCIA") or "").split("|") if x.strip()]
        if not nombres:
            nombres = ["Redondeo"]
        sap_valores = _campos(r.get("VALOR EN SAP"))
        siat_valores = _campos(r.get("VALOR EN SIAT"))
        for nombre in nombres:
            sap_val = sap_valores.get(nombre, r.get("RAZÓN SOCIAL SAP") if nombre == "Razón social" else "—")
            siat_val = siat_valores.get(nombre, r.get("RAZÓN SOCIAL SIAT") if nombre == "Razón social" else "—")
            st.markdown("<div class='ue-campo'>" + _texto(nombre) + "</div>", unsafe_allow_html=True)
            sap, siat = st.columns(2)
            sap.markdown("<div class='ue-valor'><small>SAP</small><strong>" + _texto(sap_val) + "</strong></div>", unsafe_allow_html=True)
            siat.markdown("<div class='ue-valor ue-siat'><small>SIAT</small><strong>" + _texto(siat_val) + "</strong></div>", unsafe_allow_html=True)
        if r.get("MOTIVO DE CLASIFICACIÓN MENOR"):
            st.caption("Motivo: " + str(r["MOTIVO DE CLASIFICACIÓN MENOR"]))
        if tipo == "minor":
            st.caption("La clasificación informativa no aprueba ni modifica la identidad del proveedor.")


def _explorador(filas, tipo):
    nombres = {"control":"ESTADO", "pend":"SITUACION", "minor":"TIPO DE DIFERENCIA"}
    campo_tipo = nombres[tipo]
    clave = "panel_compras_" + tipo
    bus, estado = st.columns([3, 1])
    buscar = bus.text_input("Buscar factura, NIT, proveedor, usuario o asiento",
                            key=clave+"_buscar", placeholder="Buscar factura...")
    opciones = ["Todos"] + sorted({str(x.get(campo_tipo)) for x in filas if x.get(campo_tipo)})
    tipo_elegido = estado.selectbox("Tipo de hallazgo", opciones, key=clave+"_filtro")
    if buscar.strip():
        q = buscar.strip().casefold()
        campos_busqueda = ("N.º FACTURA","FACTURA SAP","FACTURA SIAT","NIT","NIT SAP","NIT SIAT",
                           "RAZÓN SOCIAL SAP","RAZÓN SOCIAL SIAT","PROVEEDOR SAP","PROVEEDOR SIAT",
                           "USUARIO SAP","N.º ASIENTO SAP")
        filas = [f for f in filas if any(q in str(f.get(c) or "").casefold() for c in campos_busqueda)]
    if tipo_elegido != "Todos":
        filas = [f for f in filas if str(f.get(campo_tipo)) == tipo_elegido]
    filas = sorted(filas, key=lambda f: _importe(f, tipo), reverse=True)
    st.caption(f"{_cantidad(len(filas))} resultados · importes ordenados de mayor a menor")
    tamano = 15
    max_pag = max(1, (len(filas)+tamano-1)//tamano)
    pagina = st.number_input("Página", min_value=1, max_value=max_pag,
                              value=1, key=clave+"_pagina") if max_pag > 1 else 1
    visibles = filas[(pagina-1)*tamano:pagina*tamano]
    if not visibles:
        st.info("No se encontraron registros con los filtros seleccionados.")
    for f in visibles:
        factura = _numero(f, tipo)
        proveedor = (f.get("RAZÓN SOCIAL SAP") or f.get("PROVEEDOR SAP")
                     or f.get("RAZÓN SOCIAL SIAT") or f.get("PROVEEDOR SIAT") or "Proveedor no identificado")
        etiqueta = str(f.get(campo_tipo) or "").replace("CRÍTICA · ", "")
        with st.expander(f"FACTURA {factura}   ·   {_bolivianos(_importe(f,tipo))}   ·   {proveedor[:54]}   |   {etiqueta}"):
            _detalle(f, tipo)


def mostrar_panel_ejecutivo(excel_bytes: bytes, metrics: dict):
    """Añade KPIs ejecutivos, clasificación de causas y expedientes expandibles."""
    try:
        libro = load_workbook(BytesIO(excel_bytes), read_only=True, data_only=True)
        try:
            control = _filas(libro, "CONTROL GENERAL")
            pendientes = _filas(libro, "POR CONCILIAR")
            menores = _filas(libro, "DIFERENCIAS MENORES")
        finally:
            libro.close()
    except Exception as exc:
        st.warning("No fue posible preparar el panel ejecutivo: " + str(exc))
        return

    st.markdown("""<style>
    .ue-head {padding:24px 28px;background:linear-gradient(115deg,#57002b,#8b0e49);
       color:white;border-radius:14px;margin:20px 0 16px}
    .ue-head h2 {font-size:1.75rem;font-weight:850;color:white!important;margin:0 0 7px}
    .ue-head p {font-size:.88rem;color:#f5dce8;margin:0}
    .ue-card {border:1px solid #eee6eb;background:#fff;border-top:3px solid #8e0344;
       border-radius:12px;padding:14px 15px;min-height:112px}
    .ue-card small{display:block;color:#7d7580;font-size:.79rem}
    .ue-card b{display:block;color:#77013a;font-size:2rem;line-height:1.8}
    .ue-card span{display:block;color:#97929a;font-size:.7rem}
    .ue-cuf{padding:8px 11px;color:#6c5863;background:#faf7f9;border:1px solid #eee2ea;
      border-radius:7px;overflow-wrap:anywhere;font-size:.73rem;margin-bottom:15px}
    .ue-campo{color:#850a47;font-weight:850;margin:12px 0 6px;font-size:.83rem}
    .ue-valor{background:#f8f7f9;border:1px solid #ede8ee;border-radius:8px;padding:13px;min-height:75px}
    .ue-valor small{display:block;color:#928994;margin-bottom:5px;font-size:.73rem}
    .ue-valor strong{overflow-wrap:anywhere;font-size:.84rem;color:#34323a}
    .ue-siat{background:#fdf6f9}.ue-siat strong{color:#8e0d4a}
    </style>""", unsafe_allow_html=True)
    st.markdown('<div class="ue-head"><h2>CENTRO EJECUTIVO DE AUDITORÍA</h2>'
                '<p>Control de Compras SAP × SIAT · Observaciones prioritarias y consulta por factura</p></div>',
                unsafe_allow_html=True)
    cards = (("Conciliadas", metrics.get("matched",0), "Facturas cruzadas entre fuentes"),
             ("Revisión prioritaria", len(control), "Observaciones relevantes"),
             ("Por conciliar", len(pendientes), "Documentos sin correspondencia"),
             ("Informativas", len(menores), "Solo razón social o redondeos"))
    for col,(titulo,valor,desc) in zip(st.columns(4),cards):
        col.markdown(f'<div class="ue-card"><small>{titulo}</small><b>{_cantidad(valor)}</b><span>{desc}</span></div>',
                     unsafe_allow_html=True)
    st.caption(f"Sin diferencias: {_cantidad(metrics.get('ok',0))} · "
               f"SAP: {_cantidad(metrics.get('sap',0))} · SIAT: {_cantidad(metrics.get('siat',0))}")
    causas = Counter("Importes / tributación" if "IMPORTE" in str(r.get("ESTADO")) else
                     "CUF / identificadores" if "IDENTIFICADOR" in str(r.get("ESTADO")) else
                     "Fechas" if "FECHA" in str(r.get("ESTADO")) else "Revisión de identidad"
                     for r in control)
    if causas:
        st.markdown("#### Concentración de observaciones")
        for col,(nombre,n) in zip(st.columns(min(4,len(causas))),causas.most_common()):
            col.metric(nombre, _cantidad(n))
    st.markdown("#### Explorar facturas")
    t1,t2,t3=st.tabs((f"🔴 OBSERVADAS ({len(control)})",
                      f"🟠 POR CONCILIAR ({len(pendientes)})",
                      f"🟢 DIFERENCIAS MENORES ({len(menores)})"))
    with t1:
        st.caption("Abre una factura para examinar todos los campos que requieren revisión.")
        _explorador(control,"control")
    with t2:
        st.caption("Registros sin pareja en SAP o SIAT; no se aceptan conciliaciones por semejanza.")
        _explorador(pendientes,"pend")
    with t3:
        st.caption("Diferencias informativas, conservando los valores originales de ambas fuentes.")
        _explorador(menores,"minor")
